import re
import html
import mammoth

from docx import Document
"""
utlis.py >>
this is the function we are going to use >>>
1. detecting placholders <<...>> 
2. read the placeholders directly drom the uploaded word 
3. we covert the word doc into html so the format can be preserved when the email is displayed 
4. match the placeholders with the excel col names 
5. replace the placeholder with the correct value from each excel row 
6. validating the email adrees 
"""


# The actual text inside << >> is captured
# identfies wtv inside the placeholder 
PLACEHOLDER_PATTERN = r"<<\s*([A-Za-z0-9_\-\u0600-\u06FF ]+?)\s*>>"


def extract_placeholders(text):
    """
    find all the place holders inside a piece of text
    Input:
        "Dear <<name>>, welcome to <<company>>."

    Output:
        {"name", "company"}
        """

# if there is no text > there cannot be a placelder 
    if not text:
        return set()
#make sure the inpu is treated as a string 
    text = str(text)
#search the text using the placeholders patterns 
#return unique placeholder names
    matches = re.findall(
        PLACEHOLDER_PATTERN,
        text
    )

    return {
        match.strip()
        for match in matches
    }


# =========================================================
# Read placeholders DIRECTLY from Word
#
# This avoids Mammoth/HTML causing incorrect detection.
# =========================================================

def extract_docx_placeholders(file_path):

    document = Document(file_path)

    placeholders = set()

    # Check each paragraph separately
    for paragraph in document.paragraphs:
        placeholders.update(
            extract_placeholders(paragraph.text)
        )

    # Check each table cell separately
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                placeholders.update(
                    extract_placeholders(cell.text)
                )

    return placeholders

# =========================================================
# Convert Word to HTML
# Keeps Word formatting for Outlook
# =========================================================

def extract_docx_html(file_path):

    with open(file_path, "rb") as docx_file:

        result = mammoth.convert_to_html(
            docx_file
        )

    return result.value


# =========================================================
# Find matching Excel column
#
# Excel:
# name
#
# Word:
# <<name>>
# << name >>
# <<NAME>>
#
# All will match.
# =========================================================

def find_matching_column(placeholder, row):

    wanted = str(
        placeholder
    ).strip().lower()

    for column in row.index:

        column_name = str(
            column
        ).strip().lower()

        if column_name == wanted:
            return column

    return None


# =========================================================
# Replace placeholders
# Works for normal text AND Mammoth HTML
# =========================================================

def render_template_text(template, row):

    if not template:
        return ""

    result = str(template)

    # -----------------------------------------------------
    # First detect placeholders in both normal and
    # HTML-encoded forms.
    # -----------------------------------------------------

    normal_placeholders = set(
        re.findall(
            PLACEHOLDER_PATTERN,
            result
        )
    )

    encoded_pattern = (
        r"&lt;&lt;\s*"
        r"([^&<>]+?)"
        r"\s*&gt;&gt;"
    )

    encoded_placeholders = set(
        re.findall(
            encoded_pattern,
            result,
            flags=re.IGNORECASE
        )
    )

    placeholders = {
        str(p).strip()
        for p in (
            normal_placeholders
            |
            encoded_placeholders
        )
    }

    # Replace each one

    for placeholder in placeholders:

        matching_column = find_matching_column(
            placeholder,
            row
        )

        value = ""

        if matching_column is not None:

            raw_value = row[
                matching_column
            ]

            if raw_value is not None:

                value_text = str(
                    raw_value
                ).strip()

                if value_text.lower() != "nan":
                    value = value_text

        # ---------------------------------------------
        # Normal:
        # <<name>>
        # << name >>
        # ---------------------------------------------

        normal_pattern = (
            rf"<<\s*"
            rf"{re.escape(placeholder)}"
            rf"\s*>>"
        )

        result = re.sub(
            normal_pattern,
            lambda _: value,
            result,
            flags=re.IGNORECASE
        )

        # ---------------------------------------------
        # HTML/Mammoth:
        # &lt;&lt;name&gt;&gt;
        # ---------------------------------------------

        html_pattern = (
            rf"&lt;&lt;\s*"
            rf"{re.escape(placeholder)}"
            rf"\s*&gt;&gt;"
        )

        result = re.sub(
            html_pattern,
            lambda _: html.escape(value),
            result,
            flags=re.IGNORECASE
        )

    return result


# =========================================================
# Email validation
# =========================================================

def validate_email(email):

    if email is None:
        return False

    email = str(email).strip()

    if not email:
        return False

    if email.lower() == "nan":
        return False

    pattern = (
        r"^[A-Za-z0-9._%+-]+"
        r"@[A-Za-z0-9.-]+"
        r"\.[A-Za-z]{2,}$"
    )

    return bool(
        re.match(
            pattern,
            email
        )
    )